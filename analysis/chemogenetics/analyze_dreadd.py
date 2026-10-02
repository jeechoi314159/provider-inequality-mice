# -*- coding: utf-8 -*-
"""DREADD 분석: 사용 가능 소자 선별 → 시계 정합 → CNO 대 Saline 조작 검증 → 제공자 행동 정렬 PFC β."""
import os as _os, sys as _sys
_sys.path.insert(0, _os.path.abspath(_os.path.join(_os.path.dirname(_os.path.abspath(__file__)), *(['..'] * 2))))
from pvsnp_paths import *  # data locations (see pvsnp_paths.py)
import os, re, csv, glob, math, collections, statistics as st
import numpy as np
from scipy import stats
BASE = (CHEMO[:-1])
D, OUT = os.path.join(BASE, 'data'), os.path.join(BASE, 'out')
os.makedirs(OUT, exist_ok=True)
STEP = 0.125; REGION = ['BLA', 'NAc', 'PFC']
BANDS = ['theta1', 'theta2', 'beta1', 'beta2', 'gamma1', 'gamma2']
rng = np.random.default_rng(20260930)
B = list(csv.DictReader(open(os.path.join(D, 'bandpower_whole.csv'))))
rows = []

# ================= 0) 사용 가능 소자·채널 선별
# 규칙 ① 세 채널의 1/f 지수가 서로 0.03 안에 있고 sd 도 서로 2 % 안이면 합성·포화 채널로 보고 제외
#      ② 채널별 sd 중앙값이 0.01~0.5 밖이면 그 채널 제외
#      ③ 두 조건에 모두 있는 소자만 사용
def onef(R, rg):
    v = [math.log10(float(r['%s_theta1' % rg])) - math.log10(float(r['%s_gamma2' % rg]))
         for r in R if float(r['%s_theta1' % rg]) > 0 and float(r['%s_gamma2' % rg]) > 0]
    return st.median(v) if v else float('nan')
g = collections.defaultdict(list)
for r in B: g[(r['cohort'], r['condition'], int(r['device']))].append(r)
ok_ch = {}
for k, R in g.items():
    o = [onef(R, rg) for rg in REGION]
    sd = [st.median(float(r['sd_ch%d' % i]) for r in R) for i in (1, 2, 3)]
    synth = (max(o) - min(o) < 0.03) and (max(sd) - min(sd)) / max(sd) < 0.02
    for i, rg in enumerate(REGION):
        ok_ch[(k[0], k[1], k[2], rg)] = (not synth) and (0.01 < sd[i] < 0.5)
use = {}
for coh in sorted({k[0] for k in g}):
    devs = sorted({k[2] for k in g if k[0] == coh})
    for dev in devs:
        if not all((coh, c, dev) in g for c in ('CNO', 'Saline')): continue
        keep = [rg for rg in REGION if ok_ch.get((coh, 'CNO', dev, rg)) and ok_ch.get((coh, 'Saline', dev, rg))]
        if keep: use[(coh, dev)] = keep
print('0) 사용 소자·채널:')
for k in sorted(use): print('   코호트 %s dev%d → %s' % (k[0], k[1], ', '.join(use[k])))

# ================= 1) 행동 표와 시계 정합
BEH = []
import openpyxl
for f in sorted(glob.glob(os.path.join(BASE, 'raw', '**', '*timestamp_saline_CNO_condition.xlsx'), recursive=True)):
    coh = re.search(r'Group(\w)', f).group(1)
    ws = openpyxl.load_workbook(f, data_only=True)['Sheet1']
    hdr = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None: continue
        d = dict(zip(hdr, r)); d['cohort'] = coh
        d['condition'] = 'CNO' if str(d['Condition']).upper() == 'CNO' else 'Saline'
        BEH.append(d)
def num(x):
    try:
        v = float(x); return None if v != v else v
    except: return None
IX = {(r['cohort'], r['condition'], r['segment'], int(r['trial'])): float(r['seconds'])
      for r in csv.DictReader(open(os.path.join(D, 'trial_index.csv')))}
pr = [(num(b['end']), IX.get((b['cohort'], b['condition'], 'Foraging', int(b['trial'])))) for b in BEH]
pr = [(a, c) for a, c in pr if a and c]
a_ = np.array([p[0] for p in pr]); c_ = np.array([p[1] for p in pr])
print('\n1) 시계 정합: n=%d  구간길이 − 행동 end 중앙값 %+.1f s (IQR %+.1f~%+.1f)  r=%.3f'
      % (len(pr), np.median(c_-a_), *np.percentile(c_-a_, [25, 75]), np.corrcoef(a_, c_)[0, 1]))

# ================= 2) 조작 검증
def perm_within(y, x, grp, n=20000):
    y = y.copy(); x = x.copy()
    for u in np.unique(grp):
        i = grp == u; y[i] -= y[i].mean(); x[i] -= x[i].mean()
    b = (x@y)/(x@x); cnt = 0
    for _ in range(n):
        xp = x.copy()
        for u in np.unique(grp):
            i = np.where(grp == u)[0]; xp[i] = x[rng.permutation(i)]
        if abs((xp@y)/(xp@xp)) >= abs(b): cnt += 1
    return b, (cnt+1)/(n+1)
print('\n2) 조작 검증 — CNO 대 Saline, log10 전체구간 출력 (코호트 내 중심화 + 순열 20,000회)')
for seg in ['Baseline', 'Foraging']:
    print('   [%s]' % seg)
    for rg in REGION:
        line = '     %-4s' % rg
        for bn in BANDS:
            y, x, gg = [], [], []
            for r in B:
                if r['segment'] != seg: continue
                k = (r['cohort'], int(r['device']))
                if k not in use or rg not in use[k]: continue
                if float(r['artifact_frac']) > 0.01: continue
                v = float(r['%s_%s' % (rg, bn)])
                if v <= 0: continue
                y.append(math.log10(v)); x.append(1.0 if r['condition'] == 'CNO' else 0.0); gg.append(r['cohort'])
            if len(y) < 12 or len(set(x)) < 2: continue
            b, P = perm_within(np.array(y), np.array(x), np.array(gg))
            line += '  %s %+.3f(P=%.3f)' % (bn.replace('theta','θ').replace('beta','β').replace('gamma','γ'), b, P)
            rows.append(dict(analysis='manipulation', segment=seg, region=rg, band=bn, n=len(y),
                             effect=round(b, 4), P=round(P, 4)))
        print(line)
print('\n   코호트별 PFC 값 (Foraging, log10 중앙값)')
for coh in sorted({k[0] for k in use}):
    ln = '     %s' % coh
    for bn in ['beta2', 'gamma1', 'gamma2']:
        o = {}
        for cond in ['Saline', 'CNO']:
            v = [math.log10(float(r['PFC_%s' % bn])) for r in B
                 if r['cohort'] == coh and r['condition'] == cond and r['segment'] == 'Foraging'
                 and (coh, int(r['device'])) in use and 'PFC' in use[(coh, int(r['device']))]
                 and float(r['PFC_%s' % bn]) > 0 and float(r['artifact_frac']) <= 0.01]
            o[cond] = st.median(v) if v else None
        if o['Saline'] is not None and o['CNO'] is not None:
            ln += '  %s %+.2f→%+.2f (Δ%+.2f)' % (bn, o['Saline'], o['CNO'], o['CNO']-o['Saline'])
    print(ln)

# ================= 3) 제공자 행동 정렬 PFC β
index = {}
for f in sorted(glob.glob(os.path.join(D, 'spec_*_p*.npz'))):
    m = re.match(r'spec_(\w)_(CNO|Saline)_p\d+\.npz$', os.path.basename(f))
    with np.load(f, allow_pickle=True) as z:
        for k in z.files:
            if k.endswith('_band'):
                mm = re.match(r'(\w+?)_T(-?\d+)_dev(\d)_band$', k)
                index[(m.group(1), m.group(2), mm.group(1), int(mm.group(2)), int(mm.group(3)))] = (f, k)
cache = {}
def band(coh, cond, seg, tri, dev):
    k = (coh, cond, seg, tri, dev)
    if k not in index: return None
    f, key = index[k]
    if f not in cache: cache.clear(); cache[f] = np.load(f, allow_pickle=True)
    return cache[f][key]
def wm(s, t0, t1):
    i0 = max(0, int(round(t0/STEP))); i1 = min(len(s), int(round(t1/STEP)))
    if i1-i0 < 8: return np.nan
    v = s[i0:i1]; v = v[np.isfinite(v) & (v > 0)]
    return float(np.log10(v).mean()) if v.size else np.nan
WIN = {'beta2': (0.25, 2.25), 'gamma2': (-2.0, 0.0)}
BW = (-8.0, -5.0)
ev = []
for b in BEH:
    devs = [d for (c, d) in use if c == b['cohort'] and 'PFC' in use[(c, d)]]
    for align, col in (('entry', 'act'), ('retrieval', 'get')):
        t = num(b[col])
        if t is None: continue
        for dev in devs:
            a = band(b['cohort'], b['condition'], 'Foraging', int(b['trial']), dev)
            if a is None: continue
            for bn, (w0, w1) in WIN.items():
                s = a[REGION.index('PFC'), BANDS.index(bn), :]
                v = wm(s, t+w0, t+w1); bl = wm(s, t+BW[0], t+BW[1])
                if np.isfinite(v) and np.isfinite(bl):
                    ev.append(dict(cohort=b['cohort'], condition=b['condition'], align=align, band=bn, d=v-bl))
print('\n3) 제공자 행동 정렬 PFC (창 − 기준 −8~−5 s, log10)   사건 %d' % len(ev))
for align in ['entry', 'retrieval']:
    for bn in WIN:
        S = [e for e in ev if e['align'] == align and e['band'] == bn]
        cn = [e['d'] for e in S if e['condition'] == 'CNO']; sa = [e['d'] for e in S if e['condition'] == 'Saline']
        if len(cn) < 4 or len(sa) < 4: continue
        P = stats.mannwhitneyu(cn, sa, alternative='two-sided').pvalue
        per = []
        for coh in sorted({e['cohort'] for e in S}):
            c2 = [e['d'] for e in S if e['cohort'] == coh and e['condition'] == 'CNO']
            s2 = [e['d'] for e in S if e['cohort'] == coh and e['condition'] == 'Saline']
            if c2 and s2: per.append('%s %+.3f' % (coh, np.median(c2)-np.median(s2)))
        print('   %-10s %-7s Saline %+.4f (n=%d) | CNO %+.4f (n=%d)  P=%.3f   코호트별 Δ: %s'
              % (align, bn, np.median(sa), len(sa), np.median(cn), len(cn), P, ', '.join(per)))
        rows.append(dict(analysis='event_aligned', segment=align, region='PFC', band=bn, n=len(S),
                         effect=round(np.median(cn)-np.median(sa), 4), P=round(float(P), 4)))
with open(os.path.join(OUT, 'dreadd_summary.csv'), 'w', newline='') as f:
    w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
print('\nwrote out/dreadd_summary.csv')
